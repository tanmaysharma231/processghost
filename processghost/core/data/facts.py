"""Deterministic workbook facts (docs/26 flow plan): answers that code can compute exactly, so a model is not asked.

Three facts, all computed from the file with no model:

- DEPENDENCIES: which formula cells read which ranges, and, for the conditional aggregates (SUMIF, SUMIFS, COUNTIF,
  COUNTIFS, AVERAGEIF, AVERAGEIFS), which literal criteria they apply to which range. "What else reads this cell"
  is then a lookup, and "the formula only counts rows where X equals Y" is checked against the real cells.
- VALUE LOCATIONS: where a value from one row also appears on another sheet. The columns that carry the same value
  on both sides are the candidates for how rows on the two sheets are matched.
- UNIQUENESS: whether a column's values are all different (from the profile), which says how well a column could
  identify a row.

Nothing here says what a column means. It reports what is in the file; a person decides what it is for.
"""
from __future__ import annotations

import re
from typing import Dict, Iterator, List, Optional, Set, Tuple

from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries
from pydantic import BaseModel, Field

from core.data.profile import is_formula
from core.data.tools import Cell, WorkbookTools
from core.models.data import ColumnProfile, WorkbookProfile
from core.text_grounding import compact

_REF = re.compile(r"(?:(?:'([^']+)'|([A-Za-z_][A-Za-z0-9_.]*))!)?(\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)")
_CONDITIONAL = re.compile(r"\b(SUMIFS|SUMIF|COUNTIFS|COUNTIF|AVERAGEIFS|AVERAGEIF)\(", re.I)
_THOUSANDS = re.compile(r"(?<=\d)[,.](?=\d{3}\b)")


class Condition(BaseModel):
    """A literal criterion a conditional formula applies to a range, checked against the real cells."""

    sheet: str
    range: str
    criterion: str
    matching: int = 0  # cells in the range that equal the criterion
    total: int = 0  # cells in the range


class Dependency(BaseModel):
    sheet: str
    ref: str
    formula: str
    reads: List[Tuple[str, str]] = Field(default_factory=list)  # (sheet, range) the formula reads
    conditions: List[Condition] = Field(default_factory=list)


class JoinCandidate(BaseModel):
    """A pair of columns, one on each sheet, that carry the same value for this row."""

    value: str
    target_header: str
    target_column: str
    target_ref: str
    target_distinct: int
    target_rows: int
    source_header: str
    source_column: str
    source_ref: str


def norm(value: object) -> str:
    return _THOUSANDS.sub("", compact(str(value)))


def _split_args(inner: str) -> List[str]:
    args, depth, quoted, cur = [], 0, False, []
    for ch in inner:
        if ch == '"':
            quoted = not quoted
        elif not quoted and ch == "(":
            depth += 1
        elif not quoted and ch == ")":
            depth -= 1
        if ch == "," and depth == 0 and not quoted:
            args.append("".join(cur).strip())
            cur = []
        else:
            cur.append(ch)
    args.append("".join(cur).strip())
    return args


def _calls(formula: str) -> Iterator[Tuple[str, List[str]]]:
    for m in _CONDITIONAL.finditer(formula):
        depth, i, quoted = 1, m.end(), False
        while i < len(formula) and depth:
            ch = formula[i]
            if ch == '"':
                quoted = not quoted
            elif not quoted and ch == "(":
                depth += 1
            elif not quoted and ch == ")":
                depth -= 1
            i += 1
        yield m.group(1).upper(), _split_args(formula[m.end(): i - 1])


def _range_of(arg: str, own_sheet: str) -> Optional[Tuple[str, str]]:
    m = _REF.fullmatch(arg.replace("$", ""))
    if m is None:
        return None
    return (m.group(1) or m.group(2) or own_sheet), m.group(3)


def _literal(arg: str) -> Optional[str]:
    """A plain equality criterion written as a text literal; comparisons and references are not evaluated."""
    if len(arg) >= 2 and arg[0] == '"' and arg[-1] == '"':
        text = arg[1:-1].replace('""', '"')
        return None if text[:1] in "<>" else text.lstrip("=")
    return None


def cells_in(range_text: str) -> Iterator[str]:
    min_col, min_row, max_col, max_row = range_boundaries(range_text)
    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            yield "%s%d" % (get_column_letter(col), row)


def _contains(range_text: str, ref: str) -> bool:
    min_col, min_row, max_col, max_row = range_boundaries(range_text)
    row, col = coordinate_to_tuple(ref)
    return min_row <= row <= max_row and min_col <= col <= max_col


def parse_dependencies(tools: WorkbookTools) -> List[Dependency]:
    out = []
    for sheet, ref, formula in tools.formula_cells():
        reads = []
        for m in _REF.finditer(formula.replace("$", "")):
            reads.append((m.group(1) or m.group(2) or sheet, m.group(3)))
        conditions = []
        for name, args in _calls(formula):
            if name in ("SUMIFS", "AVERAGEIFS"):
                pairs = list(zip(args[1::2], args[2::2]))
            elif name == "COUNTIFS":
                pairs = list(zip(args[0::2], args[1::2]))
            else:
                pairs = [(args[0], args[1])] if len(args) >= 2 else []
            for range_arg, crit_arg in pairs:
                located, literal = _range_of(range_arg, sheet), _literal(crit_arg)
                if located and literal is not None:
                    conditions.append(Condition(sheet=located[0], range=located[1], criterion=literal))
        out.append(Dependency(sheet=sheet, ref=ref, formula=formula, reads=reads, conditions=conditions))
    return out


def _number(text: object) -> Optional[float]:
    try:
        return float(str(text).replace(",", "").strip())
    except ValueError:
        return None


def _whole_call(formula: str) -> Optional[Tuple[str, List[str]]]:
    """(name, args) if the whole formula is ONE call to a function this module can evaluate, else None."""
    m = re.match(r"^=\s*(SUMIFS|SUMIF|COUNTIFS|COUNTIF|SUM)\(", formula.strip(), re.I)
    if m is None:
        return None
    text = formula.strip()
    depth, i, quoted = 1, m.end(), False
    while i < len(text) and depth:
        ch = text[i]
        if ch == '"':
            quoted = not quoted
        elif not quoted and ch == "(":
            depth += 1
        elif not quoted and ch == ")":
            depth -= 1
        i += 1
    return (m.group(1).upper(), _split_args(text[m.end(): i - 1])) if depth == 0 and i == len(text) else None


def evaluate_formula(tools: WorkbookTools, sheet: str, formula: str) -> Optional[float]:
    """The value of a formula that is a single SUM, SUMIF, SUMIFS, COUNTIF or COUNTIFS over literal criteria, computed from the
    cells in the file. A workbook that was never calculated by a spreadsheet application has no stored results, so this is the
    only way to know what a total would show. Anything else (nested formulas, references as criteria) returns None: never a guess."""
    call = _whole_call(formula)
    if call is None:
        return None
    name, args = call

    def values(range_arg: str) -> Optional[List[Tuple[str, Optional[float]]]]:
        located = _range_of(range_arg, sheet)
        if located is None:
            return None
        out = []
        for ref in cells_in(located[1]):
            cell = tools.cell_at(located[0], ref)
            out.append((cell.value if cell is not None and cell.value is not None else "", _number(cell.value) if cell is not None and cell.value is not None else None))
        return out

    def keep(pairs: List[Tuple[str, str]]) -> Optional[List[bool]]:
        mask: Optional[List[bool]] = None
        for range_arg, crit_arg in pairs:
            vals, literal = values(range_arg), _literal(crit_arg)
            if vals is None or literal is None:
                return None
            hit = [norm(shown) == norm(literal) for shown, _n in vals]
            mask = hit if mask is None else [a and b for a, b in zip(mask, hit)]
        return mask

    if name == "SUM":
        total = 0.0
        for arg in args:
            vals = values(arg)
            if vals is None:
                number = _number(arg)
                if number is None:
                    return None
                total += number
            else:
                total += sum(n for _s, n in vals if n is not None)
        return total
    if name in ("SUMIF", "SUMIFS"):
        if name == "SUMIF":
            if len(args) < 2:
                return None
            sum_range, pairs = (args[2] if len(args) > 2 else args[0]), [(args[0], args[1])]
        else:
            sum_range, pairs = args[0], list(zip(args[1::2], args[2::2]))
        mask, vals = keep(pairs), values(sum_range)
        return None if mask is None or vals is None or len(mask) != len(vals) else sum(n for (_s, n), k in zip(vals, mask) if k and n is not None)
    pairs = [(args[0], args[1])] if name == "COUNTIF" and len(args) >= 2 else list(zip(args[0::2], args[1::2]))
    mask = keep(pairs)
    return None if mask is None else float(sum(mask))


def computed_values(tools: WorkbookTools) -> Dict[Tuple[str, str], float]:
    """(sheet, ref) -> value for every formula cell that evaluate_formula can compute."""
    out = {}
    for sheet, ref, formula in tools.formula_cells():
        value = evaluate_formula(tools, sheet, formula)
        if value is not None:
            out[(sheet, ref)] = value
    return out


def reachable_from(deps: List[Dependency], changed: List[Tuple[str, str]]) -> Set[Tuple[str, str]]:
    """Every formula cell that depends, directly or through other formulas, on any of the `changed` cells."""
    reached: Set[Tuple[str, str]] = set(changed)
    grew = True
    while grew:
        grew = False
        for d in deps:
            key = (d.sheet, d.ref)
            if key in reached:
                continue
            if any(s == rs and _contains(r, ref) for s, r in d.reads for rs, ref in list(reached) if rs == s):
                reached.add(key)
                grew = True
    return reached - set(changed)


def dependents(deps: List[Dependency], sheet: str, ref: str) -> List[Dependency]:
    """Formula cells that read `ref` on `sheet`, directly or inside a range."""
    found = []
    for d in deps:
        if any(s == sheet and _contains(r, ref) for s, r in d.reads) and not (d.sheet == sheet and d.ref == ref):
            found.append(d)
    return found


def evaluate(tools: WorkbookTools, cond: Condition) -> Condition:
    total = matching = 0
    for ref in cells_in(cond.range):
        cell = tools.cell_at(cond.sheet, ref)
        total += 1
        shown = cell.value if cell is not None and cell.value is not None else ""
        matching += int(norm(shown) == norm(cond.criterion))
    return cond.model_copy(update={"matching": matching, "total": total})


def _column_of(profile: WorkbookProfile, sheet: str, letter: str) -> Optional[ColumnProfile]:
    sh = next((s for s in profile.sheets if s.name == sheet), None)
    return next((c for c in sh.columns if c.letter == letter), None) if sh else None


def join_candidates(tools: WorkbookTools, profile: WorkbookProfile, target_sheet: str, row: int, source_sheet: str) -> List[JoinCandidate]:
    """Columns on `source_sheet` that hold the same text value as a cell in `row` of `target_sheet`, where that source column
    is unique (a row can be found by it). Numbers and dates are skipped: an amount that happens to equal another is not an
    identifier. Ordered by how well the target column identifies its own rows."""
    source_cells = {}
    for r, c, cell in tools.sheet_cells(source_sheet):
        if not is_formula(cell.formula or "") and cell.value is not None:
            source_cells.setdefault(norm(cell.value), []).append((r, c, cell))

    found = []
    for r, c, cell in tools.sheet_cells(target_sheet):
        if r != row or cell.formula is not None:
            continue
        letter = get_column_letter(c)
        target_col = _column_of(profile, target_sheet, letter)
        if target_col is None or target_col.kind not in ("text", "mixed"):
            continue
        for sr, sc, scell in source_cells.get(norm(cell.value), []):
            s_letter = get_column_letter(sc)
            source_col = _column_of(profile, source_sheet, s_letter)
            if source_col is None or source_col.kind != "text" or not source_col.unique:
                continue
            found.append(JoinCandidate(
                value=cell.value or "", target_header=target_col.header, target_column=letter, target_ref=cell.ref,
                target_distinct=target_col.distinct_count, target_rows=target_col.non_empty,
                source_header=source_col.header, source_column=s_letter, source_ref=scell.ref,
            ))
    found.sort(key=lambda j: (-(j.target_distinct / max(j.target_rows, 1)), j.target_column))
    return found
