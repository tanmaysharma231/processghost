"""Deterministic workbook profiler (docs/26 flow plan). No model, no judgment.

It states facts a reader needs to ask good questions about a workbook: which sheets exist, where the header row
is, what each column holds, which values repeat, which cells are formulas and what they read, and which layout
quirks (merged cells, empty columns, blank rows) would trip a naive reader. It does NOT say what anything means:
it never calls a column a status or a key, never compares sheets for agreement, never judges whether a value is
right. Meaning comes from the questions asked about the process, not from the data.

Formulas are read from the file as written. A workbook produced by a script (or never opened in a spreadsheet
app) has no cached results, so a formula's computed value does not exist; the profile says so
(`has_cached_values`) and reports formula text only.
"""
from __future__ import annotations

import math
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

import openpyxl
from openpyxl.utils import get_column_letter

from core import config
from core.io.storage import file_sha256
from core.models.data import ColumnProfile, FormulaCell, FormulaPattern, SheetProfile, WorkbookProfile

_REF = re.compile(r"\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?")
_HEADER_SCAN_ROWS = 10
_SAMPLES = 3


def clip(value: Any) -> str:
    """A cell as bounded text. Cell contents are data, never instructions, and never unbounded."""
    text = "" if value is None else str(value)
    return text if len(text) <= config.DATA_CELL_MAX_CHARS else text[: config.DATA_CELL_MAX_CHARS] + "..."


def is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def detect_header_row(rows: Sequence[Sequence[Any]]) -> Optional[int]:
    """1-based number of the first row that looks like a header: at least two filled cells, all of them plain text,
    and at least half as many filled cells as the widest of the first rows. Not assumed to be row 1."""
    scan = rows[:_HEADER_SCAN_ROWS]
    widest = max((sum(1 for v in r if v not in (None, "")) for r in scan), default=0)
    for i, row in enumerate(scan, start=1):
        filled = [v for v in row if v not in (None, "")]
        if len(filled) >= 2 and len(filled) >= math.ceil(widest / 2) and all(isinstance(v, str) and not is_formula(v) for v in filled):
            return i
    return None


def _kind(values: List[Any]) -> str:
    if not values:
        return "empty"
    kinds = set()
    for v in values:
        if is_formula(v):
            kinds.add("formula")
        elif isinstance(v, bool):
            kinds.add("text")
        elif isinstance(v, (int, float)):
            kinds.add("number")
        elif isinstance(v, (datetime, date)):
            kinds.add("date")
        else:
            kinds.add("text")
    return kinds.pop() if len(kinds) == 1 else "mixed"


def _row_pattern(formula: str, row: int) -> str:
    """The formula with references to its own row replaced by {r}, so a formula repeated down a column compares equal."""
    return re.sub(r"(?<![A-Za-z$])(\$?[A-Z]{1,3})%d(?!\d)" % row, r"\1{r}", formula)


def _formulas_by_column(ws) -> Tuple[List[FormulaPattern], List[FormulaCell]]:
    groups: Dict[Tuple[str, str], List[int]] = {}
    cells: Dict[Tuple[str, str], Tuple[str, str]] = {}
    for row in ws.iter_rows():
        for c in row:
            if is_formula(c.value):
                key = (c.column_letter, _row_pattern(c.value, c.row))
                groups.setdefault(key, []).append(c.row)
                cells[key] = (c.coordinate, c.value)
    patterns, singles = [], []
    for (col, pattern), rows in sorted(groups.items()):
        if len(rows) >= 2:
            patterns.append(FormulaPattern(column=col, pattern=pattern, rows=[rows[0], rows[-1]]))
        else:
            ref, formula = cells[(col, pattern)]
            singles.append(FormulaCell(ref=ref, formula=clip(formula), references=_REF.findall(formula)))
    return patterns, singles


def _profile_sheet(ws) -> SheetProfile:
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    header_row = detect_header_row(rows)
    data_start = header_row if header_row else 0  # rows[data_start:] are the rows after the header
    columns: List[ColumnProfile] = []
    empty_columns: List[str] = []
    for ci in range(1, ws.max_column + 1):
        letter = get_column_letter(ci)
        header = clip(rows[header_row - 1][ci - 1]) if header_row and ci <= len(rows[header_row - 1]) and rows[header_row - 1][ci - 1] is not None else ""
        values = [r[ci - 1] for r in rows[data_start:] if ci <= len(r) and r[ci - 1] not in (None, "")]
        if not values:
            empty_columns.append(letter)  # no data under it, whether or not it has a header
        plain = [clip(v) for v in values if not is_formula(v)]
        distinct = sorted(set(plain))
        low_cardinality = len(distinct) <= config.DATA_DISTINCT_MAX and (len(distinct) <= 5 or len(distinct) < len(plain))
        seen, samples = set(), []
        for v in plain:
            if v not in seen and len(samples) < _SAMPLES:
                seen.add(v)
                samples.append(v)
        columns.append(
            ColumnProfile(
                letter=letter, header=header, kind=_kind(values), non_empty=len(values), samples=samples,
                distinct_count=len(distinct), unique=bool(plain) and len(distinct) == len(plain),
                distinct=distinct if plain and low_cardinality else [],
            )
        )
    patterns, singles = _formulas_by_column(ws)
    last_filled = max((i for i, r in enumerate(rows, start=1) if any(v not in (None, "") for v in r)), default=0)
    return SheetProfile(
        name=ws.title, max_row=ws.max_row, max_column=ws.max_column, header_row=header_row, columns=columns,
        formula_columns=patterns, formula_cells=singles, merged_ranges=[str(r) for r in ws.merged_cells.ranges],
        empty_columns=empty_columns,
        blank_rows=[i for i, r in enumerate(rows, start=1) if i < last_filled and not any(v not in (None, "") for v in r)],
    )


def _has_cached_values(formula_wb, value_wb) -> bool:
    for ws in formula_wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if is_formula(c.value) and value_wb[ws.title][c.coordinate].value is not None:
                    return True
    return False


def profile_workbook(path: Path) -> WorkbookProfile:
    formula_wb = openpyxl.load_workbook(path, data_only=False, keep_links=False)
    value_wb = openpyxl.load_workbook(path, data_only=True, keep_links=False)
    sheets = [_profile_sheet(ws) for ws in formula_wb.worksheets]
    headers: Dict[str, set] = {}
    for sh in sheets:
        for col in sh.columns:
            if col.header:
                headers.setdefault(col.header.strip().lower(), set()).add(sh.name)
    return WorkbookProfile(
        workbook=Path(path).name, sha256=file_sha256(Path(path)), has_cached_values=_has_cached_values(formula_wb, value_wb),
        sheets=sheets, shared_headers=sorted(h for h, where in headers.items() if len(where) > 1),
    )
