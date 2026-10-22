"""Detect "one shown, many needed" (docs/26 flow plan): the recording changes one cell of a run that the process evidently covers.

A person records the task once, for one item. The process they are describing is for every item. Nothing in the speech
or the video says so, but the workbook does: the formula that reads the changed cell reads a whole run of cells (a column
range such as F18:F25), and the recording filled one of them. Where no formula gives the extent, the cells above and below
that are empty in the file, and so waiting for the same entry, give it.

The finding is raised as a question for the person, not added as a step: the extent comes from the file, and whether the same
process applies to every other item is theirs to confirm. It needs a workbook; from the screen alone "one of many" cannot be
seen.
"""
from __future__ import annotations

from typing import List, Optional, Tuple

from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries

from core.data import facts
from core.data.guard import changed_cells
from core.data.tools import WorkbookTools
from core.models.data import DataCitation, RungAttempt
from core.models.sop import Flag, Resolution, SOPDocument, SOPStep
from core.screen.cell_trail import TrailCell
from core.slots import step_window
from core.text_grounding import compact


def _cite(tools: WorkbookTools, sheet: str, ref: str) -> Optional[DataCitation]:
    cell = tools.cell_at(sheet, ref)
    if cell is None or (cell.value is None and cell.formula is None):
        return None
    return DataCitation(workbook=tools.name, sheet=sheet, ref=ref, value=cell.value, formula=cell.formula)


def _read_extent(deps: List[facts.Dependency], sheet: str, ref: str) -> Optional[Tuple[str, facts.Dependency]]:
    """The largest column range on `sheet` that holds `ref` and is read by a formula, with that formula."""
    row, col = coordinate_to_tuple(ref)
    best: Optional[Tuple[int, str, facts.Dependency]] = None
    for d in deps:
        for s, r in d.reads:
            if s != sheet or ":" not in r:
                continue
            min_col, min_row, max_col, max_row = range_boundaries(r)
            if min_col == max_col == col and min_row <= row <= max_row and (best is None or max_row - min_row > best[0]):
                best = (max_row - min_row, r, d)
    return (best[1], best[2]) if best else None


def _empty_run(tools: WorkbookTools, sheet: str, ref: str) -> Optional[str]:
    """The contiguous run of cells in the column that are empty in the file, around `ref` (the recording is filling one)."""
    row, col = coordinate_to_tuple(ref)
    letter = get_column_letter(col)

    def empty(r: int) -> bool:
        cell = tools.cell_at(sheet, "%s%d" % (letter, r))
        return cell is not None and cell.value is None and cell.formula is None

    last = tools.max_row(sheet)  # cells past the used range are not "waiting": the sheet simply ends
    top = bottom = row
    while top > 1 and empty(top - 1):
        top -= 1
    while bottom < last and empty(bottom + 1):
        bottom += 1
    return "%s%d:%s%d" % (letter, top, letter, bottom) if bottom > top else None


def add_repeat_questions(doc: SOPDocument, trail: List[TrailCell], tools: WorkbookTools) -> SOPDocument:
    changed = changed_cells(doc, trail)
    if not changed:
        return doc
    deps = facts.parse_dependencies(tools)
    asked = {compact(f.question) for f in doc.flags}
    flags, resolutions = list(doc.flags), list(doc.resolutions)
    steps = {s.step_id: s for s in doc.steps}
    by_cell = {}
    for step in doc.steps:
        fill = step.slots.get("target")
        located = [t for t in trail if fill and t.event_id in fill.event_ids]
        if located and located[-1].sheet:
            by_cell.setdefault((located[-1].sheet, located[-1].cell), step.step_id)

    for sheet, ref in dict.fromkeys(changed):
        extent, reader = _read_extent(deps, sheet, ref) or (None, None)
        basis = "the range %s!%s that %s!%s reads" % (sheet, extent, reader.sheet, reader.ref) if extent else None
        if extent is None:
            extent = _empty_run(tools, sheet, ref)
            basis = "the run of empty cells %s!%s waiting for the same entry" % (sheet, extent) if extent else None
        if extent is None:
            continue
        refs = list(facts.cells_in(extent))
        done = [r for r in refs if (sheet, r) in set(changed)]
        remaining = len(refs) - len(done)
        if remaining < 1:
            continue
        question = (
            "The recording changes %d of the %d cells in %s, %s. Is the same process repeated for the remaining %d rows?"
            % (len(done), len(refs), extent, basis, remaining)
        )
        if compact(question) in asked:
            continue
        asked.add(compact(question))
        step: Optional[SOPStep] = steps.get(by_cell.get((sheet, ref), ""))
        seg_ids, window = step_window(step) if step else ([], None)
        if not seg_ids and window is None:
            continue
        citations = [c for c in (_cite(tools, reader.sheet, reader.ref) if reader else None, _cite(tools, sheet, refs[0]), _cite(tools, sheet, ref)) if c]
        flag_id = "FLAG-%02d" % (len(flags) + 1)
        reason = "found in the workbook: %d of %d cells changed" % (len(done), len(refs))
        flags.append(Flag(
            flag_id=flag_id, family="ambiguity", question=question, severity="caution", window_segment_ids=seg_ids, window_time_range=window,
            step_id=step.step_id if step else None, status="unresolvable", origin="data", finding_kind="unresolved_requirement",
            rungs=[RungAttempt(rung="document", outcome="candidates", reason=reason, citations=citations)],
        ))
        resolutions.append(Resolution(flag_id=flag_id, verdict="unresolvable", reason=reason + "; for the person who recorded the process", citations=citations))
    return doc.model_copy(update={"flags": flags, "resolutions": resolutions})
