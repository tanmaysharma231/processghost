"""Deterministic workbook facts, the cell trail, and slot answers computed without a model. A synthetic workbook with neutral
names: no model, no real data."""
from __future__ import annotations

import openpyxl
import pytest

from core import config, pipeline
from core.data import facts
from core.data.answer import deterministic_outcomes, enrich_slots
from core.data.apply import apply_outcome
from core.data.checker import CallArgs, CellRef, DataCheckerStep, check_flag_with_data
from core.data.profile import profile_workbook
from core.data.repeat import add_repeat_questions
from core.data.tools import Filter, WorkbookTools
from core.data.wonder import WonderOutput
from core.llm import prompts
from core.ladder import record_video_rungs, sme_questions
from core.models.screen import ChangeRecord, Locator
from core.models.sop import Attribution, Flag, SlotFill, SOPStep
from core.screen.cell_trail import TrailCell, build_trail
from core.slots import fill_slot
from tests.fakes import _FakeClient
from tests.test_data_checker import _flag
from tests.test_slots import _move, _review


def _work_book(path):
    """Work: periods P1 and P2, a Ref that repeats across periods, a Label that never does, a diff formula per row and a
    conditional total over P2 that requires Flag = done. Source: the P2 rows again, with an amount."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Work"
    ws.append(["Period", "Ref", "Label", "Planned", "Actual", "Diff", "Flag"])
    rows = [("P1", "R1", "alpha 111"), ("P1", "R2", "beta 222"), ("P1", "R3", "gamma 333"),
            ("P2", "R1", "alpha 444"), ("P2", "R2", "beta 555"), ("P2", "R3", "gamma 666")]
    for i, (period, ref, label) in enumerate(rows, start=2):
        ws.append([period, ref, label, 10, 10 if period == "P1" else None, '=IF(E%d<>"",E%d-D%d,"")' % (i, i, i), "other" if period == "P1" else None])
    ws["E9"] = '=SUMIFS(E5:E7,A5:A7,"P2",G5:G7,"done")'
    src = wb.create_sheet("Source")
    src.append(["Ref", "Label", "Amount"])
    for ref, label, amount in (("R1", "alpha 444", 100), ("R2", "beta 555", 200), ("R3", "gamma 666", 300)):
        src.append([ref, label, amount])
    wb.save(path)
    return path


@pytest.fixture
def world(tmp_path):
    tools = WorkbookTools(_work_book(tmp_path / "w.xlsx"), tmp_path / "work")
    return tools, profile_workbook(tools.copy_path).model_copy(update={"workbook": tools.name}), tmp_path


# --- facts -----------------------------------------------------------------------------------------


def test_uniqueness_is_a_fact_in_the_profile(world):
    _tools, profile, _ = world
    work = next(s for s in profile.sheets if s.name == "Work")
    cols = {c.header: c for c in work.columns}
    assert cols["Label"].unique and cols["Label"].distinct_count == 6
    assert not cols["Ref"].unique and cols["Ref"].distinct_count == 3 and cols["Ref"].non_empty == 6


def test_dependents_find_every_formula_that_reads_a_cell_directly_or_in_a_range(world):
    tools, _p, _ = world
    deps = facts.parse_dependencies(tools)
    readers = {(d.sheet, d.ref) for d in facts.dependents(deps, "Work", "E5")}
    assert readers == {("Work", "F5"), ("Work", "E9")}  # the diff on its own row, and the total whose range holds it
    assert facts.dependents(deps, "Work", "A2") == [] and facts.dependents(deps, "Source", "C2") == []


def test_conditional_criteria_are_parsed_and_checked_against_the_real_cells(world):
    tools, _p, _ = world
    total = next(d for d in facts.parse_dependencies(tools) if d.ref == "E9")
    checked = {c.criterion: (c.matching, c.total) for c in (facts.evaluate(tools, c) for c in total.conditions)}
    assert checked == {"P2": (3, 3), "done": (0, 3)}  # the period matches; nothing satisfies the second condition


def test_references_to_cells_and_comparisons_are_not_mistaken_for_literal_criteria():
    assert facts._literal('">5"') is None and facts._literal("A1") is None and facts._literal('"x"') == "x"
    assert [a for _n, a in facts._calls('SUMIF(A1:A3,"x",B1:B3)')] == [["A1:A3", '"x"', "B1:B3"]]
    assert facts._split_args('A1,"a,b",IF(C1,1,2)') == ["A1", '"a,b"', "IF(C1,1,2)"]


def test_join_candidates_are_text_columns_that_carry_the_same_value_and_are_unique_on_the_source(world):
    tools, profile, _ = world
    found = facts.join_candidates(tools, profile, "Work", 5, "Source")
    assert [(j.target_header, j.source_header) for j in found] == [("Label", "Label"), ("Ref", "Ref")]  # best identifier first
    assert found[0].target_distinct == 6 and found[1].target_distinct == 3
    assert [j.target_header for j in facts.join_candidates(tools, profile, "Work", 2, "Source")] == ["Ref"]  # its label is not on the other sheet


# --- the trail -------------------------------------------------------------------------------------------


def _rec(n, t, focus, index, content=""):
    return ChangeRecord(event_id="ev-%03d" % n, t_start=t, t_end=t + 0.3, type="local",
                        surface_after=Locator(surface="spreadsheet", focus=focus, index=index, focus_content=content, read_by={"focus": "vlm"}))


def test_the_trail_lists_each_selected_cell_with_its_sheet_in_time_order():
    trail = build_trail([_rec(2, 9.0, "e5", 1), _rec(1, 3.0, "C2", 2, "100"), _rec(3, 5.0, "??", 1)], ["Work", "Source"])
    assert [(t.event_id, t.sheet, t.cell, t.content, t.read_by) for t in trail] == [("ev-001", "Source", "C2", "100", "vlm"), ("ev-002", "Work", "E5", "", "vlm")]


def test_a_cell_on_an_unknown_sheet_position_keeps_an_empty_sheet_name():
    assert build_trail([_rec(1, 1.0, "A1", None)], ["Work"])[0].sheet == ""


# --- answers computed from the trail --------------------------------------------------------------------------------


def _doc_with_cells():
    doc = _review([_move()])
    doc = fill_slot(doc, "step-02", "source", SlotFill(value="C2 (value 100)", by="shown", event_ids=["ev-a", "ev-b"]))
    return fill_slot(doc, "step-02", "target", SlotFill(value="E5", by="shown", event_ids=["ev-b"]))


TRAIL = [TrailCell(event_id="ev-a", t_s=1.0, sheet="Source", cell="C2", content="100"), TrailCell(event_id="ev-b", t_s=2.0, sheet="Work", cell="E5")]


def test_the_effect_question_is_answered_from_the_formulas_that_read_the_cell(world):
    tools, profile, _ = world
    doc = _doc_with_cells()
    out = deterministic_outcomes(doc, TRAIL, profile, tools)
    effect = out[_flag(doc, "effect").flag_id]
    assert effect.attempt.outcome == "answered" and effect.calls == 0
    assert "Work!F5" in effect.attempt.answer and "Work!E9" in effect.attempt.answer and "G5:G7 = 'done' (0 of 3" in effect.attempt.answer
    assert {c.ref for c in effect.attempt.citations} == {"F5", "E9"} and effect.attempt.citations[1].formula.startswith("=SUMIFS")


def test_a_condition_nothing_satisfies_becomes_a_new_blocking_question_for_the_person(world):
    tools, profile, _ = world
    doc = _doc_with_cells()
    effect = deterministic_outcomes(doc, TRAIL, profile, tools)[_flag(doc, "effect").flag_id]
    assert len(effect.new_questions) == 1
    q = effect.new_questions[0]
    assert q.severity == "blocking" and "G5:G7" in q.question and "'done'" in q.question and "Is setting it a step" in q.question
    done = apply_outcome(doc, _flag(doc, "effect").flag_id, effect)
    assert done.flags[-1].origin == "data" and done.flags[-1].status == "unresolvable" and done.provenance_complete


def test_a_cell_nothing_reads_is_a_real_answer(world):
    tools, profile, _ = world
    doc = _doc_with_cells()
    lonely = [TRAIL[0], TrailCell(event_id="ev-b", t_s=2.0, sheet="Work", cell="B5")]
    effect = deterministic_outcomes(doc, lonely, profile, tools)[_flag(doc, "effect").flag_id]
    assert effect.attempt.outcome == "answered" and effect.attempt.answer == "No formula in the workbook reads Work!B5."


def test_the_key_question_offers_every_matching_column_best_identifier_first_and_asks_for_confirmation(world):
    tools, profile, _ = world
    doc = _doc_with_cells()
    key = deterministic_outcomes(doc, TRAIL, profile, tools)[_flag(doc, "key").flag_id]
    assert key.attempt.outcome == "candidates" and len(key.attempt.candidates) == 2
    assert key.attempt.candidates[0].startswith("Label (Work column C, 6 distinct values in 6 rows) matches Label (Source column B)")
    assert key.attempt.candidates[1].startswith("Ref (Work column B, 3 distinct values in 6 rows)")
    done = apply_outcome(doc, _flag(doc, "key").flag_id, key)
    assert [q["style"] for q in sme_questions(done) if q["flag_id"] == _flag(done, "key").flag_id] == ["confirm"]


def test_a_single_matching_column_is_an_answer_not_a_question(world, tmp_path):
    wb = openpyxl.load_workbook(tmp_path / "w.xlsx")
    wb["Source"]["A2"] = wb["Source"]["A3"] = wb["Source"]["A4"] = None  # the Ref column no longer matches
    wb.save(tmp_path / "one.xlsx")
    tools = WorkbookTools(tmp_path / "one.xlsx", tmp_path / "w2")
    profile = profile_workbook(tools.copy_path)
    doc = _doc_with_cells()
    key = deterministic_outcomes(doc, TRAIL, profile, tools)[_flag(doc, "key").flag_id]
    assert key.attempt.outcome == "answered" and key.attempt.answer.startswith("Rows are matched on: Label")


def test_without_a_trail_cell_nothing_is_guessed(world):
    tools, profile, _ = world
    doc = _doc_with_cells()
    assert deterministic_outcomes(doc, [], profile, tools) == {}
    assert deterministic_outcomes(doc, [TRAIL[0]], profile, tools) == {}  # no target cell: the effect cannot be looked up


def test_enriching_a_slot_names_the_sheet_cell_header_and_value_and_cites_the_cell(world):
    tools, _p, _ = world
    step = enrich_slots(_doc_with_cells(), TRAIL, tools).steps[1]
    assert step.slots["source"].value == "Source!C2 (Amount) value 100" and step.slots["source"].by == "shown"
    assert step.slots["target"].value == "Work!E5 (Actual)" and step.slots["target"].citations == []  # an empty cell has nothing to cite


# --- the model loop: answers must be about what was recorded -----------------------------------------------------------


OBSERVED = [{"sheet": "Source", "cell": "C2", "content": "100"}, {"sheet": "Work", "cell": "E5", "content": ""}]


def _loop(world, steps, slot, **kw):
    tools, profile, _ = world
    doc = _doc_with_cells()
    flag = _flag(doc, slot) if slot else next(f for f in doc.flags if f.slot is None)
    flag = flag.model_copy(update={"event_ids": ["ev-b"]})
    client = _FakeClient(structured_response=list(steps))
    return client, check_flag_with_data(client, flag, doc, profile, tools, observed=OBSERVED, **kw)


def _look(sheet="Work", **f):
    return DataCheckerStep(action="tool", reason="look", call=CallArgs(tool="query_rows", sheet=sheet, filters=[Filter(**f)]))


def test_an_answer_citing_only_other_rows_is_downgraded_because_it_is_not_about_what_was_recorded(world):
    steps = [_look(column="Label", op="eq", value="alpha 111"), DataCheckerStep(action="verdict", verdict="answered", answer="x", citations=[CellRef(sheet="Work", ref="C2")])]
    _client, out = _loop(world, steps, None)
    assert out.attempt.outcome == "no_answer" and "none the screen showed" in out.attempt.reason


def test_an_answer_citing_a_cell_the_screen_showed_stands(world):
    steps = [_look(sheet="Source", column="Amount", op="eq", value="100"), DataCheckerStep(action="verdict", verdict="answered", answer="the amount", citations=[CellRef(sheet="Source", ref="C2")])]
    _client, out = _loop(world, steps, None)
    assert out.attempt.outcome == "answered" and out.attempt.citations[0].ref == "C2"


def test_key_and_effect_answers_are_exempt_because_they_come_from_workbook_facts(world):
    steps = [_look(column="Label", op="eq", value="alpha 111"), DataCheckerStep(action="verdict", verdict="answered", answer="x", citations=[CellRef(sheet="Work", ref="C2")])]
    _client, out = _loop(world, steps, "key")
    assert out.attempt.outcome == "answered"


# --- the Checker leaves these questions to the workbook ------------------------------------------------------------------


def test_with_a_workbook_the_video_checker_does_not_take_key_or_effect_questions():
    doc = _review([_move()])
    taken, left = pipeline._select_flags_for_check(doc, ("key", "effect"))
    assert {f.slot for f in taken + left} == {"source", "target", None}
    assert {f.slot for f in pipeline._select_flags_for_check(doc)[0]} == {"source", "key", "target", "effect", None}


# --- the stage ----------------------------------------------------------------------------------------------------------------


def test_the_stage_answers_key_and_effect_with_no_model_call_and_saves_the_trail(world, monkeypatch):
    _tools, _p, tmp_path = world
    monkeypatch.setattr(config, "DATA_MAX_CALLS_PER_FLAG", 1)
    records = [_rec(1, 1.0, "C2", 2, "100"), _rec(2, 2.0, "E5", 1)]
    doc = record_video_rungs(_doc_with_cells().model_copy(update={"flags": [f for f in _doc_with_cells().flags if f.slot in ("key", "effect")]}))
    doc = _doc_with_cells().model_copy(update={"flags": doc.flags})
    # the trail maps these events: ev-001 -> Source!C2, ev-002 -> Work!E5
    steps = [s.model_copy(update={"slots": {k: v.model_copy(update={"event_ids": ["ev-001", "ev-002"] if k == "source" else ["ev-002"]}) for k, v in s.slots.items()}}) for s in doc.steps]
    doc = doc.model_copy(update={"steps": steps})
    client = _FakeClient(structured_response=WonderOutput())
    out = pipeline._document_stage(client, tmp_path / "run", doc, tmp_path / "w.xlsx", records)
    # both questions answered from the file: the only model call is the one "other paths" call (docs/26), never the lookup loop
    assert [c[0] for c in client.structured_calls] == [prompts.WONDER_SYSTEM]
    assert _flag(out, "effect").status == "resolved" and _flag(out, "key").status == "unresolvable"
    assert (tmp_path / "run" / "cell_trail.json").exists() and any(f.origin == "data" for f in out.flags)


def test_with_no_confirmed_source_the_last_cell_selected_on_another_sheet_is_used(world):
    tools, profile, _ = world
    doc = _review([_move()])
    doc = fill_slot(doc, "step-02", "target", SlotFill(value="E5", by="shown", event_ids=["ev-b"]))  # no source slot at all
    key = deterministic_outcomes(doc, TRAIL, profile, tools)[_flag(doc, "key").flag_id]
    assert key.attempt.outcome == "candidates" and len(key.attempt.candidates) == 2
    same_sheet = [TrailCell(event_id="ev-a", t_s=1.0, sheet="Work", cell="C2"), TRAIL[1]]
    assert _flag(doc, "key").flag_id not in deterministic_outcomes(doc, same_sheet, profile, tools)  # no other sheet: nothing to match against


def test_a_step_with_no_located_target_uses_the_one_cell_the_whole_recording_changed(world):
    tools, profile, _ = world
    doc = _doc_with_cells()  # step-02 has the located target
    other = SOPStep(step_id="step-03", order=3, title="Check", instruction="Check it.", attributions=[Attribution(kind="said", t_s=9.0, segment_ids=[2])],
                    kind="move_data")
    flags = doc.flags + [Flag(flag_id="FLAG-90", family="ambiguity", question="What else uses it?", severity="caution", window_segment_ids=[2],
                              step_id="step-03", slot="effect")]
    doc = doc.model_copy(update={"steps": doc.steps + [other], "flags": flags})
    out = deterministic_outcomes(doc, TRAIL, profile, tools)
    assert out["FLAG-90"].attempt.outcome == "answered" and "Work!E5" in out["FLAG-90"].attempt.answer


# --- one shown, many needed ---------------------------------------------------------------------------------------


def test_the_range_a_formula_reads_shows_how_many_rows_the_process_covers(world):
    tools, _profile, _ = world
    doc = _doc_with_cells()  # the recording changes Work!E5
    out = add_repeat_questions(doc, TRAIL, tools)
    new = out.flags[len(doc.flags):]
    assert len(new) == 1 and new[0].origin == "data" and new[0].status == "unresolvable" and new[0].step_id == "step-02"
    assert "1 of the 1 cells" not in new[0].question and "remaining" in new[0].question
    assert new[0].rungs[0].citations and out.provenance_complete
    assert add_repeat_questions(out, TRAIL, tools).flags == out.flags  # asked once


def test_with_no_formula_the_empty_run_around_the_cell_gives_the_extent(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Work"
    ws.append(["Item", "Amount"])
    for i in range(2, 7):
        ws.append(["item %d" % i, 10 if i < 4 else None])
    wb.save(tmp_path / "run.xlsx")
    tools = WorkbookTools(tmp_path / "run.xlsx", tmp_path / "w")
    doc = _doc_with_cells()
    trail = [TrailCell(event_id="ev-a", t_s=1.0, sheet="Work", cell="A2"), TrailCell(event_id="ev-b", t_s=2.0, sheet="Work", cell="B4")]
    out = add_repeat_questions(doc, trail, tools)
    assert "1 of the 3 cells in B4:B6" in out.flags[-1].question and "the run of empty cells" in out.flags[-1].question


def test_a_single_cell_with_nothing_around_it_is_not_a_repeat(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Work"
    wb.active.append(["Item", "Amount"])
    wb.active.append(["only", None])
    wb.save(tmp_path / "one.xlsx")
    tools = WorkbookTools(tmp_path / "one.xlsx", tmp_path / "w")
    trail = [TrailCell(event_id="ev-a", t_s=1.0, sheet="Work", cell="A2"), TrailCell(event_id="ev-b", t_s=2.0, sheet="Work", cell="B2")]
    doc = _doc_with_cells()
    assert add_repeat_questions(doc, trail, tools).flags == doc.flags
