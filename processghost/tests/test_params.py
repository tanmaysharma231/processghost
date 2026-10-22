"""Parameters and reuse notes decided by the workbook, results drawn as results, and cell addresses written as templates.
A synthetic workbook with neutral names: no model, no real data."""
from __future__ import annotations

import openpyxl
import pytest

from core.data.params import add_parameters, find_parameters, result_notes
from core.data.profile import profile_workbook
from core.data.tools import WorkbookTools
from core.models.sop import ObservedChange, SlotFill
from core.screen.cell_trail import TrailCell
from core.slots import fill_slot
from report.flow import build_flow, render_flow_md, render_mermaid, template_phrase
from report.sop_render import render_sop_md
from tests.test_slots import _move, _review


def _book(path):
    """Work: a Month column (repeating), a Group column that repeats with no tie to any formula, an Id column whose values share a
    word with a sheet name, amounts, and a total with the month typed into it. A sheet is named for the month."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Work"
    ws.append(["Month", "Group", "Id", "Amount"])
    for i, (month, group) in enumerate([("Jan 2025", "g1"), ("Jan 2025", "g2"), ("Feb 2025", "g1"), ("Feb 2025", "g2"), ("Mar 2025", "g1"), ("Mar 2025", "g2")], start=2):
        ws.append([month, group, "WIDGET-%d" % i, None if month == "Mar 2025" else 10])
    ws["D9"] = '=SUMIFS(D6:D7,A6:A7,"Mar 2025")'
    wb.create_sheet("Source - March").append(["x"])
    wb.create_sheet("Widgets").append(["x"])
    wb.save(path)
    return path


@pytest.fixture
def world(tmp_path):
    tools = WorkbookTools(_book(tmp_path / "w.xlsx"), tmp_path / "work")
    return tools, profile_workbook(tools.copy_path)


def _doc():
    doc = _review([_move()])
    return fill_slot(doc, "step-02", "target", SlotFill(value="Work!D6 (Amount)", by="shown", event_ids=["ev-b"]))


TRAIL = [TrailCell(event_id="ev-b", t_s=2.0, sheet="Work", cell="D6")]


# --- parameters and notes -------------------------------------------------------------------------------


def test_a_value_typed_into_a_formula_is_a_parameter_with_its_evidence(world):
    tools, profile = world
    params, notes = find_parameters(_doc(), TRAIL, profile, tools)
    assert [(p.name, p.example, p.column) for p in params] == [("Month", "Mar 2025", "A")]
    assert any("Work!D9 has 'Mar 2025' typed into it" in e for e in params[0].evidence)
    assert any("Source - March" in e for e in params[0].evidence)
    assert params[0].evidence[0].startswith("Work!A holds 3 distinct values over 6 rows")


def test_the_notes_keep_the_real_names_and_say_what_would_break(world):
    tools, profile = world
    _params, notes = find_parameters(_doc(), TRAIL, profile, tools)
    kinds = {n.kind: n for n in notes}
    assert "must be changed" in kinds["hardcoded_value"].text and kinds["hardcoded_value"].citations[0].ref == "D9"
    assert "'Source - March'" in kinds["value_in_sheet_name"].text and "different sheet" in kinds["value_in_sheet_name"].text


def test_a_repeating_column_with_no_tie_and_a_sheet_name_coincidence_are_not_parameters(world):
    tools, profile = world
    params, notes = find_parameters(_doc(), TRAIL, profile, tools)
    assert "Group" not in [p.name for p in params] and "Id" not in [p.name for p in params]  # Group repeats; WIDGET-n is only a word in 'Widgets'
    assert not any("Widgets" in n.text for n in notes)


def test_each_parameter_becomes_a_question_for_the_person_once(world):
    tools, profile = world
    doc = add_parameters(_doc(), TRAIL, profile, tools)
    asked = [f for f in doc.flags if f.question.startswith("Does the Month change")]
    assert len(asked) == 1 and "'Mar 2025'" in asked[0].question and asked[0].origin == "data" and doc.provenance_complete
    assert [f.question for f in add_parameters(doc, TRAIL, profile, tools).flags].count(asked[0].question) == 1


def test_nothing_changed_means_nothing_to_say(world):
    tools, profile = world
    doc = _review([_move()])
    assert add_parameters(doc, [], profile, tools) == doc


# --- results ----------------------------------------------------------------------------------------------


def _doc_with_result(shown: bool):
    doc = _doc()
    steps = []
    for s in doc.steps:
        if s.step_id == "step-03":
            s = s.model_copy(update={"kind": "result", "observed_changes": [ObservedChange(text="the total changed", event_ids=["ev-c"], t_s=3.0)] if shown else []})
        steps.append(s)
    return doc.model_copy(update={"steps": steps})


def test_a_result_gets_the_formulas_that_produce_it_from_the_workbook(world):
    tools, _p = world
    notes = result_notes(_doc_with_result(True), TRAIL, tools)
    assert len(notes) == 1 and notes[0].step_id == "step-03" and "Work!D9" in notes[0].text and notes[0].citations[0].ref == "D9"
    assert result_notes(_doc(), TRAIL, tools) == []  # no result step: nothing to explain


def test_a_result_is_drawn_off_the_action_that_causes_it_and_labelled_by_its_evidence(world):
    tools, profile = world
    flow = build_flow(add_parameters(_doc_with_result(True), TRAIL, profile, tools))
    node = next(n for n in flow["nodes"] if n["id"] == "step-03")
    assert node["type"] == "result" and node["evidence"] == "shown" and node["event_ids"] == ["ev-c"] and "Work!D9" in node["mechanism"][0]
    results = [e for e in flow["edges"] if e["style"] == "result"]
    assert [(e["from"], e["to"]) for e in results] == [("step-02", "step-03")]
    assert all(e["to"] != "step-03" for e in flow["edges"] if e["style"] == "solid")  # the chain of actions skips the result
    assert flow["summary"]["results"] == 1 and flow["summary"]["steps"] == 2
    text = render_mermaid(flow)
    assert 'step_03(["Result: ' in text and "(shown)" in text and "step_02 -.->|result| step_03" in text


def test_a_result_nothing_on_screen_shows_says_so():
    flow = build_flow(_doc_with_result(False))
    assert next(n for n in flow["nodes"] if n["id"] == "step-03")["evidence"] == "claimed_not_shown"
    assert "claimed, not shown" in render_mermaid(flow)


def test_the_flow_and_the_report_carry_parameters_and_the_reuse_notes(world):
    tools, profile = world
    doc = add_parameters(_doc_with_result(True), TRAIL, profile, tools)
    flow = build_flow(doc)
    assert flow["parameters"][0]["name"] == "Month" and len(flow["notes"]) == 2 and flow["summary"]["parameters"] == 1
    md = render_flow_md(flow)
    assert "## Values that probably change each run" in md and "## Potential future issues" in md and "'Source - March'" in md
    report = render_sop_md(doc, video_name="v.mp4")
    assert "Reusing this flow" in report and "typed into it" in report
    assert "This is a result, not an action: nobody performs it (shown on screen)" in report and "It is produced by:" in report


# --- addresses written for reuse -----------------------------------------------------------------------------------


def test_a_cell_address_is_written_as_a_column_and_a_row_being_processed():
    assert template_phrase("Monthly Sheet!F18 (Actual)") == "Monthly Sheet, column F (Actual), the row being processed (F18 in this recording)"
    assert template_phrase("Source!C2 (Amount) value 100") == "Source, column C (Amount), the row being processed (C2 in this recording)"
    assert template_phrase("the transactions sheet") == "the transactions sheet"  # not an address: unchanged
