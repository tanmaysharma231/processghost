"""The formula evaluator, the claim guard and the workbook-versus-recording match check. A synthetic workbook with neutral
names and fake OCR text: no model, no real data."""
from __future__ import annotations

import openpyxl
import pytest

from core import pipeline
from core.data import facts
from core.data.guard import changed_cells, guard_claims, numbers_in
from core.data.match import check_match
from core.data.profile import profile_workbook
from core.data.tools import WorkbookTools
from core.ladder import record_video_rungs
from core.models.screen import ChangeRecord
from core.models.sop import Attribution, Flag, InferredClause, Resolution, SlotFill, SOPDocument, SOPStep
from core.screen.cell_trail import TrailCell
from report.sop_render import render_sop_md
from tests.fakes import _FakeClient


def _book(path):
    """Work: planned and actual amounts by period. D10 totals the P2 planned amounts, E9 totals the P2 actual amounts."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Work"
    ws.append(["Period", "Ref", "Label", "Planned", "Actual", "Flag"])
    for i, (period, ref) in enumerate([("P1", "R1"), ("P1", "R2"), ("P1", "R3"), ("P2", "R1"), ("P2", "R2"), ("P2", "R3")], start=2):
        ws.append([period, ref, "label %d" % i, 1000 * (i - 1), 100 * (i - 1) if period == "P1" else None, "done"])
    ws["E5"], ws["E6"], ws["E7"] = 100, 200, 300  # the P2 actuals are filled in: E9 = 600
    ws["E9"] = '=SUMIFS(E5:E7,A5:A7,"P2",F5:F7,"done")'
    ws["D10"] = '=SUMIF(A2:A7,"P2",D2:D7)'  # planned total for P2: 4000 + 5000 + 6000 = 15000
    ws["C11"] = '=COUNTIF(A2:A7,"P1")'
    ws["C12"] = "=SUM(D2:D4)"
    ws["C13"] = "=D2*2"  # not a single call: never evaluated
    ws["C14"] = "=SUMIF(A2:A7,B2,D2:D7)"  # a reference as the criterion: never evaluated
    wb.create_sheet("Other").append(["x"])
    wb.save(path)
    return path


@pytest.fixture
def tools(tmp_path):
    return WorkbookTools(_book(tmp_path / "w.xlsx"), tmp_path / "work")


# --- the evaluator -----------------------------------------------------------------------------------


def test_single_call_totals_are_computed_from_the_cells_in_the_file(tools):
    values = facts.computed_values(tools)
    assert values[("Work", "E9")] == 600 and values[("Work", "D10")] == 15000
    assert values[("Work", "C11")] == 3 and values[("Work", "C12")] == 6000


def test_anything_that_is_not_one_supported_call_with_literal_criteria_is_never_guessed(tools):
    values = facts.computed_values(tools)
    assert ("Work", "C13") not in values and ("Work", "C14") not in values
    assert facts.evaluate_formula(tools, "Work", "=SUM(D2:D4)+1") is None
    assert facts.evaluate_formula(tools, "Work", "=AVERAGE(D2:D4)") is None


def test_reachability_follows_formulas_that_read_a_changed_cell(tools):
    deps = facts.parse_dependencies(tools)
    assert facts.reachable_from(deps, [("Work", "E5")]) == {("Work", "E9")}
    assert facts.reachable_from(deps, [("Work", "D5")]) == {("Work", "D10"), ("Work", "C14")}  # a different column reaches different totals
    assert facts.reachable_from(deps, [("Other", "A1")]) == set()


# --- the guard -----------------------------------------------------------------------------------------


def _doc(claim, with_inserted=True):
    steps = [
        SOPStep(step_id="step-01", order=1, title="Paste", instruction="Paste it.", attributions=[Attribution(kind="said", t_s=1.0, segment_ids=[0])],
                kind="move_data", slots={"target": SlotFill(value="E5", by="shown", event_ids=["ev-1"])},
                inferred_clauses=[InferredClause(text=claim, t_s=2.0, flag_id="FLAG-01")]),
    ]
    if with_inserted:
        steps.append(SOPStep(step_id="step-ins-FLAG-02", order=2, title="Total", instruction=claim, attributions=[Attribution(kind="inferred", t_s=3.0)], flag_ids=["FLAG-02"]))
    flags = [
        Flag(flag_id="FLAG-01", family="ambiguity", question="q1", severity="caution", window_segment_ids=[0], step_id="step-01", status="resolved", event_ids=["ev-1"]),
        Flag(flag_id="FLAG-02", family="structural", question="q2", severity="caution", window_segment_ids=[0], between_step_ids=("step-01", "step-01"), status="resolved"),
    ]
    resolutions = [Resolution(flag_id=f, verdict="resolved", inferred_text=claim, event_ids=["ev-1"]) for f in ("FLAG-01", "FLAG-02")]
    return record_video_rungs(SOPDocument(run_id="r", steps=steps, flags=flags, resolutions=resolutions))


TRAIL = [TrailCell(event_id="ev-1", t_s=1.0, sheet="Work", cell="E5")]


def test_numbers_are_amounts_not_counts():
    assert numbers_in("total of 15,000 for step 2 and 600.50 or 7") == [15000.0, 600.5]


def test_the_cells_this_recording_changed_are_the_located_targets():
    assert changed_cells(_doc("x"), TRAIL) == [("Work", "E5")]


def test_a_number_that_belongs_to_a_total_over_another_column_is_withdrawn(tools):
    doc = guard_claims(_doc("The value is part of the total of 15,000 for P2.", True), TRAIL, tools)
    f1 = next(f for f in doc.flags if f.flag_id == "FLAG-01")
    res = next(r for r in doc.resolutions if r.flag_id == "FLAG-01")
    assert f1.status == "unresolvable" and res.verdict == "unresolvable" and "Work!D10" in res.reason and "does not depend" in res.reason
    assert f1.rungs[0].rung == "video" and f1.rungs[0].outcome == "no_answer" and f1.rungs[0].reason.startswith("withdrawn")
    assert [s.step_id for s in doc.steps] == ["step-01"] and doc.steps[0].inferred_clauses == [] and doc.steps[0].order == 1  # the inserted step is gone
    assert doc.provenance_complete


def test_a_number_that_belongs_to_a_cell_that_depends_on_the_change_stands(tools):
    doc = _doc("The total becomes 600.")
    assert guard_claims(doc, TRAIL, tools) == doc


def test_a_number_that_is_no_computable_cell_value_is_left_alone(tools):
    doc = _doc("Roughly 777 rows were reviewed.")
    assert guard_claims(doc, TRAIL, tools) == doc


def test_with_no_located_change_there_is_no_basis_for_the_guard(tools):
    doc = _doc("The total of 15,000.")
    assert guard_claims(doc, [], tools) == doc


# --- the match check --------------------------------------------------------------------------------------------


def _rec(total=3, n=1):
    return ChangeRecord(event_id="ev-%03d" % n, t_start=float(n), t_end=n + 0.3, type="local", chrome_after="Arial | A1 | Sheet 1 of %d | PageStyle_Work" % total)


def _check(tools, *, records=None, trail=(), texts=None):
    return check_match(records if records is not None else [_rec(2)], list(trail), texts if texts is not None else {}, tools)


def test_a_different_sheet_count_means_a_different_workbook(tools):
    m = _check(tools, records=[_rec(5)])
    assert m.verdict == "mismatch" and m.issues[0].kind == "sheet_count" and m.issues[0].fatal and "5 sheets" in m.issues[0].detail


def test_a_selected_cell_whose_content_differs_is_a_mismatch_but_an_empty_workbook_cell_is_skipped(tools):
    wrong = TrailCell(event_id="ev-1", t_s=1.0, sheet="Work", cell="C2", content="something else")
    assert _check(tools, trail=[wrong]).verdict == "mismatch"
    equal = TrailCell(event_id="ev-1", t_s=1.0, sheet="Work", cell="C2", content="label 2")
    empty = TrailCell(event_id="ev-2", t_s=2.0, sheet="Work", cell="G2", content="500")  # the workbook holds nothing here
    assert _check(tools, trail=[equal, empty]).verdict == "match"


def test_words_in_the_recording_that_the_workbook_lacks_are_a_warning_not_a_rejection(tools):
    m = _check(tools, texts={1: "Period Planned Actual label total billed Arial Format"})
    assert m.verdict == "partial" and m.issues[0].kind == "text_missing" and not m.issues[0].fatal and "billed" in m.issues[0].detail
    assert "format" not in m.issues[0].detail and "arial" not in m.issues[0].detail  # application chrome is not the document's text


def test_ocr_garble_and_sheet_names_do_not_count_as_missing(tools):
    assert _check(tools, texts={1: "Plannedd Actuall Other"}).verdict == "match"  # one stray letter, and a name from the tab strip


def test_nothing_to_compare_is_unknown_not_a_pass(tools):
    assert _check(tools, records=[]).verdict == "unknown"


# --- in the stage and the report ----------------------------------------------------------------------------------------


def test_a_workbook_that_does_not_match_is_not_used_and_says_so(tmp_path):
    path = _book(tmp_path / "w.xlsx")
    client = _FakeClient()
    doc = record_video_rungs(_doc("x", False))
    out = pipeline._document_stage(client, tmp_path / "run", doc, path, [_rec(5)])
    assert out.workbook_match.verdict == "mismatch" and client.structured_calls == []
    assert any(f["stage"] == "workbook_mismatch" for f in client.metrics.failures)
    assert all(r.rung != "document" or r.outcome == "skipped" for f in out.flags for r in f.rungs)
    assert (tmp_path / "run" / "workbook_match.json").exists()


def test_the_report_leads_with_the_match_verdict_unless_it_matches(tmp_path, tools):
    base = _doc("x", False)
    texts = {1: "billed total"}
    partial = base.model_copy(update={"workbook_match": _check(tools, texts=texts)})
    mismatch = base.model_copy(update={"workbook_match": _check(tools, records=[_rec(5)])})
    unknown = base.model_copy(update={"workbook_match": _check(tools, records=[])})
    ok = base.model_copy(update={"workbook_match": _check(tools, records=[_rec(2)])})
    assert "THE WORKBOOK MAY DIFFER FROM THE RECORDING" in render_sop_md(partial, video_name="v.mp4")
    assert "THE WORKBOOK DOES NOT MATCH THE RECORDING" in render_sop_md(mismatch, video_name="v.mp4")
    assert "not compared with the recording" in render_sop_md(unknown, video_name="v.mp4")
    assert "WORKBOOK" not in render_sop_md(ok, video_name="v.mp4").upper().replace("THE WORKBOOK WAS", "")
