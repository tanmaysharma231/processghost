"""Reading edited cells reliably: the formula-bar parse, a workbook saved after the recording, the second look at the pixels,
and the guard on a claim that names another row's value. Synthetic workbook, fake OCR and fake client: no model."""
from __future__ import annotations

import cv2
import numpy as np
import openpyxl
import pytest

from core import config
from core.data.guard import guard_claims
from core.data.match import Disputed, check_match
from core.data.tools import WorkbookTools
from core.models.screen import ChangeRecord, Locator, ScreenEvent, ScreenTimeline
from core.screen.cell_trail import TrailCell
from core.screen.surfaces.spreadsheet import SpreadsheetProfile
from core.screen.verify_reads import VerifyOutput, Choice, make_confirm
from tests.fakes import _FakeClient
from tests.test_data_guard import _doc


# --- the formula-bar parse -----------------------------------------------------------------------------------------------


@pytest.mark.parametrize("chrome, ref, content", [
    ("H3 | fx2.= | ReadyforPayment | Find | Sheet 1of3", "H3", "ReadyforPayment"),  # OCR read the icons beside the glyph as "2.="
    ("J2 | fxX | ab | Find", "J2", "ab"),  # edit mode shows a cancel icon next to the glyph
    ("G2 | Ifx, = | =F2-E2 | Find", "G2", "=F2-E2"),  # a stray letter before the glyph
    ("I2 | fx | Pending | Find", "I2", "Pending"),
    ("C7 | fx | | Find | FindAll", "C7", ""),  # an empty cell stays empty: never the toolbar text that follows
])
def test_the_content_is_read_past_glyph_debris(chrome, ref, content):
    loc = SpreadsheetProfile().locate(chrome)
    assert (loc.focus, loc.focus_content) == (ref, content)


# --- the match check -----------------------------------------------------------------------------------------------------


def _book(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Work"
    ws.append(["Item", "Result", "Who"])
    ws.append(["A", "Accepted", "abcd"])  # the workbook is saved AFTER the recording changed row 2
    ws.append(["B", "Open", None])
    wb.create_sheet("Other").append(["x"])
    wb.save(path)
    return path


@pytest.fixture
def tools(tmp_path):
    return WorkbookTools(_book(tmp_path / "w.xlsx"), tmp_path / "work")


def _rec(event_id, focus, content, side="after"):
    loc = Locator(surface="spreadsheet", index=1, focus=focus, focus_content=content)
    kw = {"surface_after": loc} if side == "after" else {"surface_before": loc}
    return ChangeRecord(event_id=event_id, t_start=1.0, t_end=1.3, type="local", chrome_after="x | Sheet 1 of 2 | PageStyle_Work", **kw)


def _trail(event_id, cell, content, t=1.0):
    return TrailCell(event_id=event_id, t_s=t, sheet_index=1, sheet="Work", cell=cell, content=content)


def test_a_cell_the_recording_changed_may_hold_the_other_state_in_the_workbook(tools):
    trail = [_trail("ev-1", "B2", "Pending")]  # the recording showed the before-state; the file is the end state
    m = check_match([_rec("ev-1", "B2", "Pending")], trail, {}, tools, changed={("Work", "B2")})
    assert m.verdict == "partial" and not any(i.fatal for i in m.issues) and "saved from before or after" in m.issues[0].detail


def test_every_reading_of_a_cell_counts_not_only_the_first(tools):
    trail = [_trail("ev-1", "B2", "Pending", 1.0), _trail("ev-2", "B2", "Accepted", 2.0)]
    assert check_match([_rec("ev-1", "B2", "Pending"), _rec("ev-2", "B2", "Accepted")], trail, {}, tools).verdict == "match"
    before_side = [_rec("ev-1", "B2", "Accepted", side="before")]  # the before-side reading of a record is a reading too
    assert check_match(before_side, [_trail("ev-1", "B2", "Pending")], {}, tools).verdict == "match"


def test_an_unchanged_cell_that_disagrees_is_fatal_unless_the_pixels_back_the_workbook(tools):
    trail = [_trail("ev-1", "B3", "Closed")]
    records = [_rec("ev-1", "B3", "Closed")]
    assert check_match(records, trail, {}, tools).verdict == "mismatch"  # no second look available: held against the workbook

    def answer(choice):
        return lambda disputed: {(d.sheet, d.cell): choice for d in disputed}

    assert check_match(records, trail, {}, tools, confirm=answer("read")).verdict == "mismatch"
    backed = check_match(records, trail, {}, tools, confirm=answer("workbook"))
    assert backed.verdict == "partial" and "supports the workbook" in backed.issues[0].detail
    unclear = check_match(records, trail, {}, tools, confirm=answer("neither"))
    assert unclear.verdict == "partial" and "could not settle it" in unclear.issues[0].detail
    assert check_match(records, trail, {}, tools, confirm=lambda disputed: {}).verdict == "mismatch"  # an unanswered cell stays against it


def test_the_formula_bar_of_a_formula_cell_shows_the_formula_not_its_value(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Work"
    ws.append(["Item", "Qty", "Double"])
    ws.append(["A", 4, "=B2*2"])  # written by a script: no stored result, the formula is all the file holds
    wb.create_sheet("Other").append(["x"])
    wb.save(tmp_path / "f.xlsx")
    tools = WorkbookTools(tmp_path / "f.xlsx", tmp_path / "fwork")
    same = check_match([_rec("ev-1", "C2", "=B2*2")], [_trail("ev-1", "C2", "=B2*2")], {}, tools)
    assert same.verdict == "match"
    other = check_match([_rec("ev-1", "C2", "=B2*3")], [_trail("ev-1", "C2", "=B2*3")], {}, tools)
    assert other.verdict == "mismatch" and "'=B2*2'" in other.issues[0].detail


def test_the_second_look_is_asked_once_for_all_disputed_cells(tools):
    asked = []

    def confirm(disputed):
        asked.append(list(disputed))
        return {}

    trail = [_trail("ev-1", "B3", "Closed"), _trail("ev-2", "C3", "Someone")]
    check_match([_rec("ev-1", "B3", "Closed"), _rec("ev-2", "C3", "Someone")], trail, {}, tools, confirm=confirm)
    # C3 is empty in the workbook, so nothing there can disagree: only B3 is asked about, and in a single call
    assert len(asked) == 1 and [(d.cell, d.read, d.workbook) for d in asked[0]] == [("B3", "Closed", "Open")]


# --- the second look -----------------------------------------------------------------------------------------------------


def _timeline_with_frame(tmp_path):
    frames = tmp_path / "tl"
    (frames / "keyframes").mkdir(parents=True)
    img = np.full((60, 400, 3), 255, np.uint8)
    cv2.imwrite(str(frames / "keyframes" / "a.png"), img)
    cv2.imwrite(str(frames / "keyframes" / "b.png"), img)
    ev = ScreenEvent(event_id="ev-1", t_start=1.0, t_end=1.3, type="local", changed_pct=1.0, coverage_pct=1.0,
                     before_frame="keyframes/b.png", after_frame="keyframes/a.png")
    return ScreenTimeline(video="v", duration_s=3.0, width=400, height=60, events=[ev], edit_band=(10, 40)), frames


def _disputed():
    return [Disputed("Work", "B3", "Closed", "Open", "ev-1", "after")]


def test_a_choice_of_the_workbook_stands_only_if_local_ocr_also_shows_its_text(tmp_path):
    timeline, frames = _timeline_with_frame(tmp_path)
    reply = VerifyOutput(answers=[Choice(id="ev-1/after", choice="workbook")])
    agrees = make_confirm(_FakeClient(structured_response=reply), timeline, frames, tmp_path / "v", lambda img: "fx | Open")
    assert agrees(_disputed()) == {("Work", "B3"): "workbook"}
    unsupported = make_confirm(_FakeClient(structured_response=reply), timeline, frames, tmp_path / "v2", lambda img: "fx | Closed")
    assert unsupported(_disputed()) == {("Work", "B3"): "neither"}  # the model sided with the file but the pixels do not show it
    no_ocr = make_confirm(_FakeClient(structured_response=reply), timeline, frames, tmp_path / "v3", None)
    assert no_ocr(_disputed()) == {("Work", "B3"): "neither"}


def test_a_choice_of_the_reading_needs_no_corroboration_and_the_prompt_shows_both_candidates(tmp_path):
    timeline, frames = _timeline_with_frame(tmp_path)
    client = _FakeClient(structured_response=VerifyOutput(answers=[Choice(id="ev-1/after", choice="read")]))
    assert make_confirm(client, timeline, frames, tmp_path / "v", lambda img: "")(_disputed()) == {("Work", "B3"): "read"}
    system, user = client.structured_calls[0][0], client.structured_calls[0][1]
    assert "Candidate read: 'Closed'" in user and "Candidate workbook: 'Open'" in user
    assert "only evidence" in system and "not hints" in system


def test_a_failed_second_look_confirms_nothing_and_is_recorded(tmp_path):
    timeline, frames = _timeline_with_frame(tmp_path)
    client = _FakeClient(structured_response=RuntimeError("model down"))
    assert make_confirm(client, timeline, frames, tmp_path / "v", None)(_disputed()) == {}
    assert any(f["stage"] == "verify_read_failed" for f in client.metrics.failures)


def test_over_the_crop_cap_the_rest_are_not_sent(tmp_path, monkeypatch):
    timeline, frames = _timeline_with_frame(tmp_path)
    monkeypatch.setattr(config, "VLM_READ_MAX_CROPS", 1)
    many = [Disputed("Work", "B%d" % i, "x", "y", "ev-1", "after") for i in (3, 4, 5)]
    client = _FakeClient(structured_response=VerifyOutput(answers=[Choice(id="ev-1/after", choice="read")]))
    make_confirm(client, timeline, frames, tmp_path / "v", None)(many)
    assert client.structured_calls[0][1].count("image ") == 1


# --- a claim that names another row's value --------------------------------------------------------------------------------


def _name_book(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Work"
    ws.append(["Item", "Who"])
    for i, who in enumerate(["Ann Lee", "Bo Chan", None, "Ann Lee"], start=2):
        ws.append(["item %d" % i, who])
    wb.save(path)
    return path


@pytest.fixture
def named(tmp_path):
    return WorkbookTools(_name_book(tmp_path / "n.xlsx"), tmp_path / "nwork")


NAME_TRAIL = [TrailCell(event_id="ev-1", t_s=1.0, sheet="Work", cell="B4", content="typed text")]


def test_a_claim_that_names_another_value_of_the_changed_column_is_withdrawn(named):
    doc = guard_claims(_doc("The invoice was approved by Bo Chan.", False), NAME_TRAIL, named)
    flag = next(f for f in doc.flags if f.flag_id == "FLAG-01")
    res = next(r for r in doc.resolutions if r.flag_id == "FLAG-01")
    assert flag.status == "unresolvable" and "another value of column B" in res.reason and "typed text" in res.reason
    assert doc.steps[0].inferred_clauses == [] and doc.provenance_complete


def test_a_claim_that_also_holds_what_the_cell_read_stands(named):
    doc = _doc("It now reads typed text, unlike Bo Chan's row.", False)
    assert guard_claims(doc, NAME_TRAIL, named) == doc


def test_with_no_settled_reading_or_another_column_the_guard_stays_silent(named):
    doc = _doc("The invoice was approved by Bo Chan.", False)
    unread = [TrailCell(event_id="ev-1", t_s=1.0, sheet="Work", cell="B4", content="")]
    assert guard_claims(doc, unread, named) == doc
    elsewhere = [TrailCell(event_id="ev-1", t_s=1.0, sheet="Work", cell="A4", content="typed text")]  # column A: no claim about the Who column
    assert guard_claims(doc, elsewhere, named) == doc
