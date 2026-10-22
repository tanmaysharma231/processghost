"""Unit tests for core/io/documents.py -- no API key, no network call, no LLM anywhere in this
module: source ingestion is deterministic per docs/19-pre-meeting-intake.md Phase 1.

Fixture workbooks are written at test time with openpyxl into tmp_path -- never a checked-in
scenario file, per the repo's scenario-agnostic discipline (tests/test_prompts_scenario_agnostic.py
enforces the same principle for prompts). One additional test at the bottom is gated on the real
data/ files existing and asserts only the real scenario's sheet names, which is where a genuine
scenario assertion belongs.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from core.io.documents import load_source
from core.text_grounding import quote_is_grounded


def _write_text(tmp_path: Path, content: str) -> Path:
    p = tmp_path / "brief.txt"
    p.write_text(content, encoding="utf-8")
    return p


def _write_workbook(tmp_path: Path, *, repeated_formula_rows: int = 5) -> Path:
    """A small fixture workbook exercising the same shapes as the real scenario: multiple sheets,
    a header row, a column of identical relative formulas, and one distinct total-row formula."""
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Reference List"
    ws1.append(["ID", "Name"])
    ws1.append(["R-001", "Alpha"])

    ws2 = wb.create_sheet("Working Sheet")
    ws2.append(["ID", "Expected", "Actual", "Variance"])
    for i in range(2, 2 + repeated_formula_rows):
        ws2.append([f"R-{i:03d}", 10, "", ""])
        ws2[f"D{i}"] = f"=IF(C{i}<>\"\",C{i}-B{i},\"\")"
    ws2[f"B{2 + repeated_formula_rows}"] = f"=SUM(B2:B{1 + repeated_formula_rows})"

    p = tmp_path / "workbook.xlsx"
    wb.save(p)
    return p


# --- text ingestion -----------------------------------------------------------

def test_text_source_splits_into_one_passage_per_paragraph(tmp_path):
    path = _write_text(tmp_path, "Para one line a.\nPara one line b.\n\nPara two.\n")

    doc = load_source(path, source_id="SRC-01")

    assert doc.kind == "text"
    assert len(doc.passages) == 2
    assert doc.passages[0].locator == "lines 1-2"
    assert doc.passages[1].locator == "lines 4-4"


def test_every_passage_is_groundable_in_the_source_file(tmp_path):
    content = "First fact stated here.\n\nSecond fact, stated differently.\n"
    path = _write_text(tmp_path, content)

    doc = load_source(path, source_id="SRC-01")

    assert doc.passages
    for passage in doc.passages:
        assert quote_is_grounded(passage.text, content)  # a passage is a quote of the file, not a paraphrase


def test_unsupported_suffix_raises_rather_than_producing_an_empty_source(tmp_path):
    path = tmp_path / "brief.pdf"
    path.write_text("irrelevant", encoding="utf-8")

    with pytest.raises(ValueError):
        load_source(path, source_id="SRC-01")


def test_non_ascii_source_text_round_trips_intact(tmp_path):
    path = _write_text(tmp_path, "Amount due is ₹500 for this line item.\n")

    doc = load_source(path, source_id="SRC-01")

    assert "₹500" in doc.passages[0].text


# --- spreadsheet ingestion -----------------------------------------------------------

def test_spreadsheet_source_names_every_sheet(tmp_path):
    path = _write_workbook(tmp_path)

    doc = load_source(path, source_id="SRC-01")

    locators = [p.locator for p in doc.passages]
    assert any(loc.startswith("Reference List!") for loc in locators)
    assert any(loc.startswith("Working Sheet!") for loc in locators)
    header_text = next(p.text for p in doc.passages if p.locator == "Working Sheet!1:1")
    assert "Working Sheet" in header_text


def test_header_passage_carries_column_letters_not_just_names(tmp_path):
    path = _write_workbook(tmp_path)

    doc = load_source(path, source_id="SRC-01")

    header_text = next(p.text for p in doc.passages if p.locator == "Working Sheet!1:1")
    assert "D=Variance" in header_text


def test_repeated_relative_formula_collapses_to_one_range_passage(tmp_path):
    path = _write_workbook(tmp_path, repeated_formula_rows=24)

    doc = load_source(path, source_id="SRC-01")

    formula_passages = [p for p in doc.passages if "formula" in p.text]
    range_passages = [p for p in formula_passages if ":" in p.locator.split("!", 1)[1]]
    assert len(range_passages) == 1
    assert range_passages[0].locator == "Working Sheet!D2:D25"


def test_distinct_total_formula_gets_its_own_single_cell_passage(tmp_path):
    path = _write_workbook(tmp_path)

    doc = load_source(path, source_id="SRC-01")

    total_passages = [p for p in doc.passages if "SUM(" in p.text]
    assert len(total_passages) == 1
    assert ":" not in total_passages[0].locator.split("!", 1)[1]


def test_spreadsheet_passage_count_stays_small_not_one_per_cell(tmp_path):
    path = _write_workbook(tmp_path, repeated_formula_rows=24)

    doc = load_source(path, source_id="SRC-01")

    assert len(doc.passages) < 10


def test_every_spreadsheet_passage_points_at_a_real_sheet(tmp_path):
    path = _write_workbook(tmp_path)

    doc = load_source(path, source_id="SRC-01")

    sheets = set(openpyxl.load_workbook(path).sheetnames)
    assert doc.passages
    for passage in doc.passages:
        assert passage.locator.split("!", 1)[0] in sheets


def test_row_cap_truncation_is_recorded_not_silent(tmp_path, monkeypatch):
    from core import config

    monkeypatch.setattr(config, "INTAKE_MAX_ROWS_PER_SHEET", 3)
    path = _write_workbook(tmp_path, repeated_formula_rows=10)

    doc = load_source(path, source_id="SRC-01")

    assert "truncated" in doc.ingest_note


# --- the real scenario, gated on the real files existing -----------------------------------------

_REAL_WORKBOOK = Path(__file__).parent.parent.parent / "data" / "ProcessGhost_Bills_Tracker_v4_final.xlsx"

pytestmark_real = pytest.mark.skipif(not _REAL_WORKBOOK.exists(), reason="real scenario workbook not present")


@pytestmark_real
def test_real_workbook_yields_all_three_documented_sheet_names():
    doc = load_source(_REAL_WORKBOOK, source_id="SRC-01")

    sheet_prefixes = {p.locator.split("!", 1)[0] for p in doc.passages}
    assert {"Master Bills", "Monthly Actuals", "Transactions - March"} <= sheet_prefixes
